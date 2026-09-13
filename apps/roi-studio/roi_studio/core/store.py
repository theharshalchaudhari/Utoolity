"""The annotation store: one canonical row schema and every output file.

Rows are kept in ONE shape (CANON_COLUMNS) and projected to the chosen view
only at write time, which is what stops a format toggle from producing a
ragged mixed-schema spreadsheet.

Deliberately free of Qt so the whole persistence path is testable headlessly.
"""

from __future__ import annotations

import csv
import json
import os
import shutil
import tempfile
from datetime import datetime

from ..config import (APP_VERSION, CANON_COLUMNS, COLUMN_WIDTHS, FULL_VIEW,
                      JSON_NAME, MAP_NAME, NDIGITS, ROW_TYPES, TEXT_COLUMNS,
                      XLSX_NAME)
from . import geometry as geo
from .io_safe import (WriteReport, atomic_write, timestamped_sibling,
                      write_text_atomic)

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    HAS_XLSX = True
except Exception:                                    # pragma: no cover
    Workbook = load_workbook = None
    Font = PatternFill = Alignment = get_column_letter = None
    HAS_XLSX = False

LEGACY_CSV_NAMES = ("roi_annotations.csv",)

SPREADSHEET_EXTS = (".xlsx", ".xlsm", ".xltx", ".xltm")


def load_workbook_rows(path):
    """Read a workbook as a list of row tuples.

    openpyxl validates the *file name*, not the content, so it refuses a
    rolling backup called `roi_annotations.xlsx.bak` even though the bytes
    are a perfectly good workbook.  Anything without a supported extension is
    copied to a correctly named temporary file first."""
    if not HAS_XLSX:
        return []
    path = str(path)
    temporary = ""
    try:
        if not path.lower().endswith(SPREADSHEET_EXTS):
            fd, temporary = tempfile.mkstemp(suffix=".xlsx")
            os.close(fd)
            shutil.copyfile(path, temporary)
            path = temporary
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            return list(wb.active.iter_rows(values_only=True))
        finally:
            wb.close()
    finally:
        if temporary:
            try:
                os.remove(temporary)
            except OSError:
                pass


def rows_to_dicts(rows):
    """Turn raw workbook rows into dicts keyed by the header line."""
    if not rows:
        return []
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    out = []
    for values in rows[1:]:
        if values is None or all(v is None or str(v).strip() == ""
                                 for v in values):
            continue
        item = {}
        for i, key in enumerate(header):
            if not key:
                continue
            item[key] = "" if i >= len(values) or values[i] is None \
                else str(values[i])
        out.append(item)
    return out


# ══════════════════════════════════════════════════════════════
# JSON LAYOUT HELPERS
# ══════════════════════════════════════════════════════════════
def _compact_map_block(mapping, item_indent="  ", close_indent=""):
    """Render {"KEY": [[[x, y], ...]], ...} with ONE LINE PER KEY.

    json.dump(indent=2) puts every single number on its own line, which makes
    a 500-camera map unreadable.  Keys are still json.dumps()-escaped, so this
    stays valid JSON, and every file is parsed back before it replaces the
    old one anyway."""
    if not mapping:
        return "{}"
    parts = ["%s%s: %s" % (item_indent, json.dumps(k),
                           json.dumps(v, separators=(", ", ", ")))
             for k, v in mapping.items()]
    return "{\n" + ",\n".join(parts) + "\n" + close_indent + "}"


def dump_map_json(roi_map) -> str:
    """The map file: one roi_key per line."""
    return _compact_map_block(roi_map, "    ", "") + "\n"


def dump_full_json(full) -> str:
    """Metadata pretty-printed, coordinate maps compact."""
    lines = ["{"]
    skip = ("rois", "no_roi", "sources", "shapes")
    for key in [k for k in full if k not in skip]:
        lines.append('  %s: %s,' % (json.dumps(key), json.dumps(full[key])))
    lines.append('  "rois": %s,' % _compact_map_block(full["rois"], "    ", "  "))
    lines.append('  "shapes": %s,' % _compact_map_block(full["shapes"], "    ", "  "))
    lines.append('  "no_roi": %s,'
                 % json.dumps(full["no_roi"], separators=(", ", ", ")))
    lines.append('  "sources": %s' % _compact_map_block(full["sources"],
                                                        "    ", "  "))
    lines.append("}")
    return "\n".join(lines) + "\n"


# ══════════════════════════════════════════════════════════════
class AnnotationStore:
    """Owns every row and every output file for one batch folder."""

    def __init__(self):
        self.folder = ""
        self.rows = []
        self.xlsx_path = ""
        self.json_path = ""
        self.map_path = ""
        self.stamp = None            # mtime of the spreadsheet we last saw

    # ── setup ─────────────────────────────────────────────
    def bind(self, folder) -> None:
        self.folder = str(folder)
        self.xlsx_path = os.path.join(self.folder, XLSX_NAME)
        self.json_path = os.path.join(self.folder, JSON_NAME)
        self.map_path = os.path.join(self.folder, MAP_NAME)
        self.rows = []

    # ── canonical row construction ────────────────────────
    @staticmethod
    def make_row(fname, row_type, polys_px, norm_polys, comment,
                 width=0, height=0, kinds=None):
        site_id, cam_token = geo.extract_site_cam(fname)
        return {
            "image_name": fname,
            "roi_key": geo.make_roi_key(fname),
            "site_id": site_id,
            "cam_number": cam_token,
            "pixel_coords": geo.fmt_polys(polys_px),
            "normalized_coords": geo.fmt_polys(norm_polys),
            "total_polygons": len(polys_px) if row_type == "roi" else 0,
            "image_width": int(width or 0),
            "image_height": int(height or 0),
            "row_type": row_type if row_type in ROW_TYPES else "roi",
            "comment": comment or "",
            "shape_types": geo.fmt_shape_types(kinds) if row_type == "roi" else "",
        }

    @staticmethod
    def canonicalize(raw):
        """Bring any legacy row up to the current schema.

        Older builds wrote either the site layout (roi_coordinate) or the full
        layout (pixel_coords / normalized_coords), and toggling between them
        left rows carrying both.  Everything folds back into one shape here."""
        row = {k: "" for k in CANON_COLUMNS}
        for key in CANON_COLUMNS:
            if key in raw and raw[key] is not None:
                row[key] = str(raw[key]).strip()

        fname = row["image_name"]
        if not fname:
            # some very old files put the bare basename in site_id
            fname = str(raw.get("site_id", "") or "").strip()
            row["image_name"] = fname

        if not row["normalized_coords"]:
            legacy = raw.get("roi_coordinate", "")
            if legacy:
                polys = geo.parse_multi_polys(legacy)
                if polys:
                    row["normalized_coords"] = geo.fmt_polys(polys)

        # re-render both coordinate cells in the current bracket format
        for key in ("pixel_coords", "normalized_coords"):
            if row[key]:
                row[key] = geo.fmt_polys(geo.parse_multi_polys(row[key]))

        if fname:
            site_id, cam_token = geo.extract_site_cam(fname)
            row["site_id"] = row["site_id"] or site_id
            row["cam_number"] = row["cam_number"] or cam_token
            row["roi_key"] = row["roi_key"] or geo.make_roi_key(fname)

        rtype = row["row_type"] or "roi"
        row["row_type"] = rtype if rtype in ROW_TYPES else "roi"

        count = len(geo.parse_multi_polys(row["pixel_coords"]
                                          or row["normalized_coords"]))
        try:
            row["total_polygons"] = int(float(row["total_polygons"] or 0))
        except (TypeError, ValueError):
            row["total_polygons"] = count
        for key in ("image_width", "image_height"):
            try:
                row[key] = int(float(row[key] or 0))
            except (TypeError, ValueError):
                row[key] = 0

        # shape_types is padded to the polygon count, so a file written
        # before the column existed describes plain polygons
        row["shape_types"] = geo.fmt_shape_types(
            geo.parse_shape_types(row["shape_types"], count)) if count else ""

        # The JSON export is built from normalized_coords.  A legacy row that
        # only has pixels would otherwise vanish from the JSON, so derive them
        # whenever the image size is recoverable.
        if not row["normalized_coords"] and row["pixel_coords"] \
                and row["image_width"] and row["image_height"]:
            px = geo.parse_multi_polys(row["pixel_coords"])
            if px:
                row["normalized_coords"] = geo.fmt_polys(
                    geo.polys_to_norm(px, row["image_width"], row["image_height"]))
        return row

    # ── loading ───────────────────────────────────────────
    def load(self):
        """Prefer the xlsx, fall back to its .bak, then to a legacy csv.

        Returns a list of human-readable notes."""
        notes = []
        candidates = [(self._load_xlsx, self.xlsx_path, XLSX_NAME),
                      (self._load_xlsx, self.xlsx_path + ".bak", XLSX_NAME + ".bak")]
        for name in LEGACY_CSV_NAMES:
            path = os.path.join(self.folder, name)
            candidates.append((self._load_csv, path, name))
            candidates.append((self._load_csv, path + ".bak", name + ".bak"))

        for loader, path, label in candidates:
            if not path or not os.path.isfile(path):
                continue
            try:
                raw_rows = loader(path)
            except Exception as exc:
                notes.append("Could not read %s (%s)" % (label, exc))
                continue
            if raw_rows:
                self.rows = [self.canonicalize(r) for r in raw_rows]
                self.rows = [r for r in self.rows if r["image_name"]]
                notes.append("Loaded %d existing row(s) from %s"
                             % (len(self.rows), label))
                self.touch_stamp()
                return notes
        self.touch_stamp()
        return notes

    # ── external change detection ─────────────────────────
    def touch_stamp(self) -> None:
        """Remember the spreadsheet as it is on disk right now."""
        try:
            self.stamp = os.path.getmtime(self.xlsx_path) \
                if os.path.isfile(self.xlsx_path) else None
        except OSError:
            self.stamp = None

    def changed_externally(self) -> bool:
        """True when the spreadsheet on disk is not the one we last wrote.

        Someone else - another session, a sync client, a person editing in
        Excel - has been here since, so overwriting it would discard their
        work silently."""
        if not self.xlsx_path:
            return False
        try:
            exists = os.path.isfile(self.xlsx_path)
        except OSError:
            return False
        if not exists:
            return False
        if self.stamp is None:
            return True
        try:
            return abs(os.path.getmtime(self.xlsx_path) - self.stamp) > 1.0
        except OSError:
            return False

    @staticmethod
    def _load_csv(path):
        with open(path, "r", encoding="utf-8-sig", newline="") as fh:
            return [{k: ("" if v is None else v) for k, v in row.items()
                     if k is not None} for row in csv.DictReader(fh)]

    @staticmethod
    def _load_xlsx(path):
        return rows_to_dicts(load_workbook_rows(path))

    # ── row helpers ───────────────────────────────────────
    def purge(self, fname) -> None:
        self.rows = [r for r in self.rows if r["image_name"] != fname]

    def add(self, row) -> None:
        self.rows.append(row)

    def rows_for(self, fname):
        return [r for r in self.rows if r["image_name"] == fname]

    def row_for(self, fname):
        rows = self.rows_for(fname)
        return rows[0] if rows else None

    def image_names(self):
        seen, out = set(), []
        for row in self.rows:
            name = row["image_name"]
            if name not in seen:
                seen.add(name)
                out.append(name)
        return out

    # ── projection ────────────────────────────────────────
    @staticmethod
    def project(row, view):
        out = {}
        for key in view:
            if key == "roi_coordinate":
                out[key] = row.get("normalized_coords", "")
            else:
                out[key] = row.get(key, "")
        return out

    # ── JSON build ────────────────────────────────────────
    def build_json(self, rows=None):
        """Merge every image's polygons under its roi_key.

        Two images that share a site+camera (different timestamps) have their
        polygon lists concatenated, in natural filename order.  A key that has
        at least one ROI never appears in the no_roi list."""
        rows = self.rows if rows is None else rows
        roi_map, shape_map, sources = {}, {}, {}
        no_roi_keys, warnings = [], []

        ordered = sorted(rows, key=lambda r: geo.natural_key(r["image_name"]))
        for row in ordered:
            key = str(row.get("roi_key", "")).strip()
            fname = row.get("image_name", "")
            if not key:
                warnings.append("no roi_key for %s" % fname)
                continue
            rtype = row.get("row_type", "")
            if rtype == "roi":
                polys = geo.parse_multi_polys(row.get("normalized_coords", ""))
                if not polys:
                    # An ROI row that cannot be normalised would vanish from
                    # the JSON.  Say so rather than quietly shipping a short
                    # file - re-saving the image repairs it.
                    if row.get("pixel_coords"):
                        warnings.append(
                            "%s has pixel coords but no image size - "
                            "re-save it to include it in the JSON" % fname)
                    continue
                if key in roi_map:
                    warnings.append("merged %s into %s" % (fname, key))
                roi_map.setdefault(key, []).extend(polys)
                shape_map.setdefault(key, []).extend(
                    geo.parse_shape_types(row.get("shape_types", ""), len(polys)))
                sources.setdefault(key, []).append(fname)
            elif rtype == "no_roi":
                if key not in no_roi_keys:
                    no_roi_keys.append(key)
                sources.setdefault(key, []).append(fname)

        no_roi_keys = [k for k in no_roi_keys if k not in roi_map]

        rounded = {}
        for key in sorted(roi_map, key=geo.natural_key):
            rounded[key] = [[[round(float(x), NDIGITS), round(float(y), NDIGITS)]
                             for x, y in poly] for poly in roi_map[key]]

        full = {
            "generated_at": datetime.now().isoformat(timespec="seconds"),
            "generator": "ROI Studio v%s" % APP_VERSION,
            "source_folder": self.folder,
            "coordinate_space": "normalized",
            "decimals": NDIGITS,
            "roi_count": len(rounded),
            "no_roi_count": len(no_roi_keys),
            "total_polygons": sum(len(v) for v in rounded.values()),
            "rois": rounded,
            "shapes": {k: shape_map.get(k, []) for k in sorted(rounded, key=geo.natural_key)},
            "no_roi": sorted(no_roi_keys, key=geo.natural_key),
            "sources": {k: sources.get(k, []) for k in sorted(sources, key=geo.natural_key)},
        }
        return full, rounded, warnings

    # ── writing ───────────────────────────────────────────
    def flush(self, view_columns=None, rows=None) -> WriteReport:
        """Write the xlsx, the full json and the map json.

        Every file goes through atomic_write and is read back to prove it
        parses before the old copy is replaced."""
        report = WriteReport()
        if not self.folder:
            report.errors.append("no output folder")
            return report
        if not os.path.isdir(self.folder):
            report.errors.append("output folder no longer exists")
            return report

        rows = self.rows if rows is None else rows
        view_columns = list(view_columns or FULL_VIEW)

        full, roi_map, warns = self.build_json(rows)
        report.warnings.extend(warns[:3])

        self._write_json(self.json_path, dump_full_json(full), report)
        self._write_json(self.map_path, dump_map_json(roi_map), report)

        if HAS_XLSX:
            self._write_xlsx(view_columns, rows, report)
        else:
            report.errors.append(
                "openpyxl is not installed, so no spreadsheet can be written")
        return report

    def _write_json(self, path, text, report) -> None:
        ok, err = write_text_atomic(path, text, verify_json=True)
        if ok:
            report.written.append(path)
        else:
            report.errors.append("%s: %s" % (os.path.basename(path), err))

    def _write_xlsx(self, view_columns, rows, report) -> None:
        projected = [self.project(r, view_columns) for r in rows]

        def write(tmp):
            wb = Workbook()
            ws = wb.active
            ws.title = "ROI"
            ws.append(list(view_columns))
            for row in projected:
                ws.append([row.get(c, "") for c in view_columns])

            head_font = Font(bold=True, color="FFFFFF")
            head_fill = PatternFill("solid", fgColor="2F6BD8")
            for idx, name in enumerate(view_columns, start=1):
                cell = ws.cell(row=1, column=idx)
                cell.font = head_font
                cell.fill = head_fill
                cell.alignment = Alignment(horizontal="left", vertical="center")
                ws.column_dimensions[get_column_letter(idx)].width = \
                    COLUMN_WIDTHS.get(name, 16)
                # Force TEXT on every string column so Excel can never
                # re-interpret a coordinate string as numbers or dates.
                if name in TEXT_COLUMNS:
                    for r in range(2, len(projected) + 2):
                        ws.cell(row=r, column=idx).number_format = "@"
            ws.freeze_panes = "A2"
            if projected:
                ws.auto_filter.ref = "A1:%s%d" % (
                    get_column_letter(len(view_columns)), len(projected) + 1)
            wb.save(tmp)

        def verify(tmp):
            wb = load_workbook(tmp, read_only=True)
            try:
                if wb.active.max_row < 1:
                    raise ValueError("workbook came back empty")
            finally:
                wb.close()

        ok, err = atomic_write(self.xlsx_path, write, verify)
        if ok:
            report.written.append(self.xlsx_path)
            self.touch_stamp()
            return
        # The usual cause is the file being open in Excel; keep the work by
        # writing a timestamped sibling rather than losing the save.
        fallback = timestamped_sibling(self.xlsx_path)
        ok2, err2 = atomic_write(fallback, write, verify, keep_backup=False)
        if ok2:
            self.xlsx_path = fallback
            report.warnings.append("%s locked, wrote %s"
                                   % (XLSX_NAME, os.path.basename(fallback)))
            report.written.append(fallback)
            self.touch_stamp()
        else:
            report.errors.append("xlsx: %s / %s" % (err, err2))
